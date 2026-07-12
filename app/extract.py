"""Best-effort text extraction from uploaded documents.

Supports PDF, DOCX, XLSX and plain-text formats. Any failure returns an empty
string rather than raising, so an unparseable upload never breaks the flow.
"""

from __future__ import annotations

import io

_MAX_CHARS = 20000  # keep extracted context bounded for prompts/storage


def extract_text(data: bytes, filename: str, content_type: str = "") -> str:
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf") or "pdf" in content_type:
            text = _from_pdf(data)
        elif name.endswith(".docx") or "wordprocessingml" in content_type:
            text = _from_docx(data)
        elif name.endswith(".xlsx") or "spreadsheetml" in content_type:
            text = _from_xlsx(data)
        elif name.endswith((".txt", ".md", ".csv", ".json", ".log")):
            text = data.decode("utf-8", errors="replace")
        else:
            # Unknown type: try utf-8, otherwise give up gracefully.
            text = data.decode("utf-8", errors="ignore")
    except Exception:
        return ""
    text = (text or "").strip()
    return text[:_MAX_CHARS]


def _from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _from_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    return "\n".join(p.text for p in document.paragraphs)


def _from_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for ws in wb.worksheets:
        lines.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)
